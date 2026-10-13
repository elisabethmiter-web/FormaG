"""Email campaigns: each company keeps a client list and can email all of it (or a group) at once.

Contacts come from three places: everyone a packet is sent to, CSV/Excel imports, and people added by hand.
Managers see the whole company list. Staff see the contacts they added plus the clients of their own packets,
and can only email those.

Every campaign email carries the company's name and mailing address and a one-click unsubscribe link,
as Canada's anti-spam law (CASL) and the big mailbox providers require. Unsubscribed people are never
emailed again, even if they are imported again.
"""
import csv
import io
import json
import os
import re
import secrets
import threading
import time
import uuid
from datetime import datetime, timedelta, timezone
from html import escape

from flask import Response, abort, flash, g, redirect, render_template, request, url_for

DAILY_LIMIT = int(os.environ.get("CAMPAIGN_DAILY_LIMIT", "2000"))   # per company, protects your sending reputation
PER_MINUTE = max(1, int(os.environ.get("CAMPAIGN_PER_MINUTE", "120")))
MAX_IMPORT_ROWS = 20000
EMAIL_RE = re.compile(r"^[^@\s,;<>\"']+@[^@\s,;<>\"']+\.[A-Za-z]{2,}$")
WORKER = uuid.uuid4().hex[:12]
STALE_SECONDS = 90

SCHEMA = """
CREATE TABLE IF NOT EXISTS contacts (
  id INTEGER PRIMARY KEY,
  company_id INTEGER NOT NULL,
  name TEXT DEFAULT '',
  email TEXT NOT NULL COLLATE NOCASE,
  organization TEXT DEFAULT '',
  tags TEXT DEFAULT '',
  source TEXT NOT NULL DEFAULT 'manual',
  owner_id INTEGER,
  unsubscribed_at TEXT,
  created_at TEXT NOT NULL,
  UNIQUE (company_id, email)
);
CREATE TABLE IF NOT EXISTS campaigns (
  id INTEGER PRIMARY KEY,
  company_id INTEGER NOT NULL,
  subject TEXT NOT NULL DEFAULT '',
  body TEXT NOT NULL DEFAULT '',
  button_text TEXT DEFAULT '',
  button_url TEXT DEFAULT '',
  audience TEXT NOT NULL DEFAULT 'all',
  status TEXT NOT NULL DEFAULT 'draft',
  created_by_id INTEGER,
  created_by TEXT DEFAULT '',
  created_at TEXT NOT NULL,
  sent_at TEXT,
  finished_at TEXT,
  base_url TEXT DEFAULT '',
  worker TEXT,
  heartbeat_at TEXT
);
CREATE TABLE IF NOT EXISTS campaign_recipients (
  id INTEGER PRIMARY KEY,
  campaign_id INTEGER NOT NULL REFERENCES campaigns(id),
  contact_id INTEGER,
  email TEXT NOT NULL,
  name TEXT DEFAULT '',
  organization TEXT DEFAULT '',
  token TEXT UNIQUE NOT NULL,
  status TEXT NOT NULL DEFAULT 'queued',
  error TEXT DEFAULT '',
  sent_at TEXT,
  opened_at TEXT
);
CREATE INDEX IF NOT EXISTS ix_contacts_company ON contacts(company_id);
CREATE INDEX IF NOT EXISTS ix_campaigns_company ON campaigns(company_id);
CREATE INDEX IF NOT EXISTS ix_recipients_campaign ON campaign_recipients(campaign_id, status);
"""


def init_db(conn):
    """Create the tables and add everyone already sent a packet to their company's client list."""
    conn.executescript(SCHEMA)
    cols = {r[1] for r in conn.execute("PRAGMA table_info(companies)")}
    if "mailing_address" not in cols:
        conn.execute("ALTER TABLE companies ADD COLUMN mailing_address TEXT DEFAULT ''")
    conn.execute(
        "INSERT OR IGNORE INTO contacts (company_id, name, email, source, owner_id, created_at) "
        "SELECT company_id, client_name, LOWER(TRIM(client_email)), 'packet', created_by_id, MIN(created_at) "
        "FROM packets WHERE company_id IS NOT NULL AND client_email LIKE '%_@_%' "
        "GROUP BY company_id, LOWER(TRIM(client_email))")


def clean_tags(raw):
    seen = []
    for t in re.split(r"[,;|]", raw or ""):
        t = re.sub(r"\s+", " ", t).strip()[:40]
        if t and t.lower() not in [s.lower() for s in seen]:
            seen.append(t)
    return ", ".join(seen[:12])


def tag_list(tags):
    return [t.strip() for t in (tags or "").split(",") if t.strip()]


def remember_contact(conn, company_id, name, email, owner_id, source="packet"):
    """Add a client to the company list (called when a packet is created). Never changes an existing contact."""
    email = (email or "").strip().lower()
    if EMAIL_RE.match(email):
        conn.execute("INSERT OR IGNORE INTO contacts (company_id, name, email, source, owner_id, created_at) "
                     "VALUES (?,?,?,?,?,?)", (company_id, (name or "").strip()[:120], email, source, owner_id,
                                              datetime.now(timezone.utc).replace(microsecond=0).isoformat()))


def text_to_html(text):
    """The campaign body is plain text: blank lines start paragraphs and web addresses become links."""
    out = []
    for para in re.split(r"\n\s*\n", (text or "").strip()):
        h = escape(para.strip()).replace("\n", "<br>")
        h = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", h)
        h = re.sub(r"(https?://[^\s<]+[^\s<.,;:!?)])",
                   r'<a href="\1" style="color:#24408e">\1</a>', h)
        if h:
            out.append(f'<p style="margin:0 0 16px">{h}</p>')
    return "".join(out)


def personalize(text, r):
    name = (r.get("name") or "").strip()
    first = name.split()[0] if name else ""
    return (text.replace("{first_name}", first or "there").replace("{name}", name or "there")
                .replace("{organization}", r.get("organization") or ""))


def render_email(camp, company, r, unsub_url, pixel_url=None):
    """(html, text) for one recipient."""
    body = personalize(camp["body"], r)
    address = (company.get("mailing_address") or "").strip()
    btn_html = btn_text = ""
    if camp["button_text"] and camp["button_url"]:
        btn_html = (f'<p style="margin:8px 0 24px"><a href="{escape(camp["button_url"])}" style="display:inline-block;'
                    f'background:#24408e;color:#ffffff;padding:12px 22px;border-radius:6px;text-decoration:none;'
                    f'font-weight:bold">{escape(camp["button_text"])}</a></p>')
        btn_text = f"\n{camp['button_text']}: {camp['button_url']}\n"
    pixel = f'<img src="{escape(pixel_url)}" width="1" height="1" alt="" style="display:block;border:0">' if pixel_url else ""
    footer = (f'{escape(company["name"])}' + (f"<br>{escape(address).replace(chr(10), '<br>')}" if address else "")
              + (f'<br>Contact: <a href="mailto:{escape(company["notify_email"])}" style="color:#5b6475">'
                 f'{escape(company["notify_email"])}</a>' if company.get("notify_email") else "")
              + f'<br><br>You are receiving this email because you are a client or contact of {escape(company["name"])}. '
                f'<a href="{escape(unsub_url)}" style="color:#5b6475">Unsubscribe</a>')
    html = f"""<!doctype html><html><body style="margin:0;padding:0;background:#f6f7f9">
<div style="max-width:600px;margin:0 auto;padding:24px 16px;font-family:Arial,Helvetica,sans-serif;color:#1d2433;line-height:1.55;font-size:15px">
<div style="background:#ffffff;border:1px solid #dfe3ea;border-radius:8px;padding:28px 28px 12px">
<p style="margin:0 0 20px;font-size:13px;font-weight:bold;letter-spacing:.06em;text-transform:uppercase;color:#5b6475">{escape(company["name"])}</p>
{text_to_html(body)}{btn_html}</div>
<p style="font-size:12px;color:#5b6475;line-height:1.5;margin:16px 4px 0">{footer}</p>{pixel}
</div></body></html>"""
    text = (f"{body.strip()}\n{btn_text}\n--\n{company['name']}\n" + (f"{address}\n" if address else "")
            + (f"Contact: {company['notify_email']}\n" if company.get("notify_email") else "")
            + f"Unsubscribe: {unsub_url}\n")
    return html, text


# ---------------------------------------------------------------- reading client lists from files
HEADER_KEYS = {
    "email": ("email", "e-mail", "courriel", "email address", "mail"),
    "first": ("first", "first name", "firstname", "given name", "prénom", "prenom"),
    "last": ("last", "last name", "lastname", "surname", "family name", "nom de famille"),
    "name": ("name", "full name", "contact", "contact name", "client", "client name", "customer", "customer name", "nom"),
    "organization": ("company", "organization", "organisation", "business", "company name", "account", "entreprise"),
    "tags": ("tags", "tag", "group", "groups", "list", "category", "segment", "type"),
}


def read_rows(f):
    """Rows (lists of strings) from an uploaded CSV or Excel file."""
    filename = (f.filename or "").lower()
    data = f.read()
    if filename.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001
            raise ValueError("That Excel file couldn't be opened. Save it again as .xlsx or CSV and retry.")
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else str(v).strip() for v in row])
            if len(rows) > MAX_IMPORT_ROWS + 1:
                break
        return rows
    if filename.endswith(".xls"):
        raise ValueError("Old .xls files aren't supported. In Excel choose File → Save As → .xlsx or CSV.")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)][:MAX_IMPORT_ROWS + 1]


def map_columns(rows):
    """Work out which column holds what. Returns (column map, rows without the header)."""
    if not rows:
        return {}, []
    head = [h.lower().strip().replace("_", " ") for h in rows[0]]
    cols = {}
    for key, names in HEADER_KEYS.items():
        for i, h in enumerate(head):
            if i in cols.values():
                continue
            if h in names or (key == "email" and "email" in h) or (key == "email" and "e-mail" in h):
                cols[key] = i
                break
    if "email" in cols:
        return cols, rows[1:]
    # No header row: find the column that holds email addresses; the first other column is the name.
    for i, v in enumerate(rows[0]):
        if EMAIL_RE.match(v.strip()):
            other = [j for j in range(len(rows[0])) if j != i]
            return ({"email": i, "name": other[0]} if other else {"email": i}), rows
    return {}, rows


# ---------------------------------------------------------------- routes
def register(A):
    app, db = A.app, A.db

    def now():
        return A.now()

    def contact_scope(alias="c"):
        """Contacts the signed-in user may see and email."""
        if A.is_manager():
            return f"{alias}.company_id = ?", [A.cid()]
        pcond, pargs = A.packet_scope()
        return (f"{alias}.company_id = ? AND ({alias}.owner_id = ? OR {alias}.email IN "
                f"(SELECT LOWER(TRIM(client_email)) FROM packets WHERE {pcond}))", [A.cid(), g.user["id"]] + pargs)

    def campaign_scope():
        if A.is_manager():
            return "company_id = ?", [A.cid()]
        return "company_id = ? AND created_by_id = ?", [A.cid(), g.user["id"]]

    def load_campaign(camp_id):
        cond, args = campaign_scope()
        return db().execute(f"SELECT * FROM campaigns WHERE id=? AND {cond}", [camp_id] + args).fetchone() or abort(404)

    def tags_in_scope():
        cond, args = contact_scope()
        found = {}
        for (t,) in db().execute(f"SELECT tags FROM contacts c WHERE {cond} AND tags != ''", args):
            for tag in tag_list(t):
                found.setdefault(tag.lower(), tag)
        return sorted(found.values(), key=str.lower)

    def audience_sql(audience):
        cond, args = contact_scope()
        sql = f"FROM contacts c WHERE {cond} AND c.unsubscribed_at IS NULL"
        if audience.startswith("tag:"):
            sql += " AND (',' || LOWER(REPLACE(c.tags, ', ', ',')) || ',') LIKE ?"
            args = args + [f"%,{audience[4:].lower()},%"]
        return sql, args

    def audience_count(audience):
        sql, args = audience_sql(audience)
        return db().execute(f"SELECT COUNT(*) {sql}", args).fetchone()[0]

    def audience_label(audience):
        return f"Clients tagged “{audience[4:]}”" if audience.startswith("tag:") else "Everyone on the client list"

    def sent_last_24h():
        since = (datetime.now(timezone.utc) - timedelta(hours=24)).replace(microsecond=0).isoformat()
        return db().execute("SELECT COUNT(*) FROM campaign_recipients r JOIN campaigns k ON k.id = r.campaign_id "
                            "WHERE k.company_id=? AND r.status='sent' AND r.sent_at >= ?", (A.cid(), since)).fetchone()[0]

    def stats(camp_id):
        row = db().execute("SELECT COUNT(*) total, SUM(status='sent') sent, SUM(status='failed') failed, "
                           "SUM(status='queued') queued, SUM(status='skipped') skipped, "
                           "SUM(opened_at IS NOT NULL) opened FROM campaign_recipients WHERE campaign_id=?",
                           (camp_id,)).fetchone()
        return {k: row[k] or 0 for k in row.keys()}

    app.jinja_env.globals.update(audience_label=audience_label, tag_list=tag_list)

    # ------------------------------------------------ public: unsubscribe + open tracking (no login)
    @app.route("/u/<token>", methods=["GET", "POST"])
    def unsubscribe(token):
        r = db().execute("SELECT r.*, k.company_id FROM campaign_recipients r JOIN campaigns k ON k.id = r.campaign_id "
                         "WHERE r.token=?", (token,)).fetchone() or abort(404)
        g.company = A.get_company(r["company_id"])
        if request.method == "POST":  # the button on the page, or the mail app's one-click unsubscribe
            db().execute("UPDATE contacts SET unsubscribed_at=COALESCE(unsubscribed_at, ?) WHERE company_id=? AND email=?",
                         (now(), r["company_id"], r["email"]))
            db().execute("UPDATE campaign_recipients SET status='skipped', error='unsubscribed' WHERE status='queued' "
                         "AND email=? AND campaign_id IN (SELECT id FROM campaigns WHERE company_id=?)",
                         (r["email"], r["company_id"]))
            db().commit()
            if request.form.get("List-Unsubscribe") == "One-Click":
                return "", 200
            return render_template("unsubscribe.html", email=r["email"], done=True, token=token)
        done = db().execute("SELECT unsubscribed_at FROM contacts WHERE company_id=? AND email=?",
                            (r["company_id"], r["email"])).fetchone()
        return render_template("unsubscribe.html", email=r["email"], done=bool(done and done[0]), token=token)

    PIXEL = (b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!\xf9\x04\x01\x00\x00\x00\x00"
             b",\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;")

    @app.route("/o/<token>.gif")
    def email_opened(token):
        db().execute("UPDATE campaign_recipients SET opened_at=COALESCE(opened_at, ?) WHERE token=?", (now(), token))
        db().commit()
        return Response(PIXEL, mimetype="image/gif", headers={"Cache-Control": "no-store"})

    # ------------------------------------------------ client list
    @app.route("/contacts")
    @A.login_required
    def contacts():
        q = request.args.get("q", "").strip()[:100]
        tag = request.args.get("tag", "").strip()[:40]
        show = request.args.get("show", "all")
        cond, args = contact_scope()
        sql = f"FROM contacts c WHERE {cond}"
        if q:
            like = f"%{q.replace('%', '').replace('_', '')}%"
            sql += " AND (c.name LIKE ? OR c.email LIKE ? OR c.organization LIKE ?)"
            args = args + [like, like, like]
        if tag:
            sql += " AND (',' || LOWER(REPLACE(c.tags, ', ', ',')) || ',') LIKE ?"
            args = args + [f"%,{tag.lower()},%"]
        if show == "unsubscribed":
            sql += " AND c.unsubscribed_at IS NOT NULL"
        elif show == "subscribed":
            sql += " AND c.unsubscribed_at IS NULL"
        total = db().execute(f"SELECT COUNT(*) {sql}", args).fetchone()[0]
        rows = db().execute(f"SELECT c.* {sql} ORDER BY c.created_at DESC, c.id DESC LIMIT 300", args).fetchall()
        cond0, args0 = contact_scope()
        counts = db().execute(f"SELECT COUNT(*), SUM(c.unsubscribed_at IS NULL) FROM contacts c WHERE {cond0}",
                              args0).fetchone()
        return render_template("contacts.html", contacts=rows, total=total, q=q, tag=tag, show=show,
                               tags=tags_in_scope(), all_count=counts[0] or 0, active_count=counts[1] or 0)

    @app.route("/contacts/add", methods=["POST"])
    @A.login_required
    def contact_add():
        email = request.form.get("email", "").strip().lower()[:160]
        if not EMAIL_RE.match(email):
            flash("Enter a valid email address.", "error")
        elif request.form.get("consent") != "on":
            flash("Tick the box to confirm this person agreed to hear from you.", "error")
        else:
            cur = db().execute("INSERT OR IGNORE INTO contacts (company_id, name, email, organization, tags, source, "
                               "owner_id, created_at) VALUES (?,?,?,?,?,?,?,?)",
                               (A.cid(), request.form.get("name", "").strip()[:120], email,
                                request.form.get("organization", "").strip()[:120],
                                clean_tags(request.form.get("tags", "")), "manual", g.user["id"], now()))
            db().commit()
            flash(f"{email} added to your client list." if cur.rowcount else
                  f"{email} is already on your company's client list.", "ok" if cur.rowcount else "error")
        return redirect(url_for("contacts"))

    @app.route("/contacts/import", methods=["POST"])
    @A.login_required
    def contact_import():
        f = request.files.get("file")
        if not f or not f.filename:
            flash("Choose a CSV or Excel file to import.", "error")
            return redirect(url_for("contacts"))
        if request.form.get("consent") != "on":
            flash("Tick the box to confirm the people on this list agreed to hear from you.", "error")
            return redirect(url_for("contacts"))
        try:
            cols, rows = map_columns(read_rows(f))
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("contacts"))
        if "email" not in cols:
            flash("No email column found. Make sure the first row has a heading like “Email”.", "error")
            return redirect(url_for("contacts"))
        extra_tags = clean_tags(request.form.get("tags", ""))

        def cell(row, key):
            i = cols.get(key)
            return row[i].strip() if i is not None and i < len(row) and row[i] else ""

        added = updated = invalid = dupes = unsub = 0
        seen = set()
        for row in rows[:MAX_IMPORT_ROWS]:
            email = cell(row, "email").lower().strip("<>").strip()
            if not email:
                continue
            if not EMAIL_RE.match(email):
                invalid += 1
                continue
            if email in seen:
                dupes += 1
                continue
            seen.add(email)
            name = cell(row, "name") or " ".join(x for x in (cell(row, "first"), cell(row, "last")) if x)
            tags = clean_tags(", ".join(x for x in (cell(row, "tags"), extra_tags) if x))
            org = cell(row, "organization")
            old = db().execute("SELECT * FROM contacts WHERE company_id=? AND email=?", (A.cid(), email)).fetchone()
            if old is None:
                db().execute("INSERT INTO contacts (company_id, name, email, organization, tags, source, owner_id, "
                             "created_at) VALUES (?,?,?,?,?,?,?,?)",
                             (A.cid(), name[:120], email, org[:120], tags, "import", g.user["id"], now()))
                added += 1
            else:
                if old["unsubscribed_at"]:
                    unsub += 1   # someone who unsubscribed is never signed up again by an import
                db().execute("UPDATE contacts SET name=CASE WHEN name='' THEN ? ELSE name END, "
                             "organization=CASE WHEN organization='' THEN ? ELSE organization END, tags=? WHERE id=?",
                             (name[:120], org[:120], clean_tags(old["tags"] + ", " + tags), old["id"]))
                updated += 1
        db().commit()
        parts = [f"{added} added", f"{updated} already on the list (details filled in)"]
        if invalid:
            parts.append(f"{invalid} skipped (not a valid email)")
        if dupes:
            parts.append(f"{dupes} duplicate row{'s' if dupes != 1 else ''} skipped")
        if unsub:
            parts.append(f"{unsub} had unsubscribed and stay unsubscribed")
        flash("Import finished: " + ", ".join(parts) + ".", "ok")
        return redirect(url_for("contacts"))

    @app.route("/contacts/<int:contact_id>", methods=["POST"])
    @A.login_required
    def contact_update(contact_id):
        cond, args = contact_scope()
        c = db().execute(f"SELECT * FROM contacts c WHERE c.id=? AND {cond}", [contact_id] + args).fetchone() or abort(404)
        action = request.form.get("action")
        if action == "edit":
            db().execute("UPDATE contacts SET name=?, organization=?, tags=? WHERE id=?",
                         (request.form.get("name", "").strip()[:120], request.form.get("organization", "").strip()[:120],
                          clean_tags(request.form.get("tags", "")), contact_id))
            flash("Contact saved.", "ok")
        elif action == "unsubscribe":
            db().execute("UPDATE contacts SET unsubscribed_at=? WHERE id=? AND unsubscribed_at IS NULL", (now(), contact_id))
            flash(f"{c['email']} won't receive campaign emails.", "ok")
        elif action == "resubscribe":
            if request.form.get("confirm") != "on":
                flash("Only resubscribe someone who asked to be added back. Tick the box to confirm.", "error")
            else:
                db().execute("UPDATE contacts SET unsubscribed_at=NULL WHERE id=?", (contact_id,))
                flash(f"{c['email']} will receive campaign emails again.", "ok")
        elif action == "delete":
            if c["unsubscribed_at"]:
                flash("Unsubscribed contacts are kept so they're never emailed again by mistake.", "error")
            else:
                db().execute("DELETE FROM contacts WHERE id=?", (contact_id,))
                flash(f"{c['email']} removed from the client list.", "ok")
        db().commit()
        return redirect(request.referrer if (request.referrer or "").startswith(request.host_url) else url_for("contacts"))

    @app.route("/contacts/export.csv")
    @A.login_required
    def contacts_export():
        cond, args = contact_scope()
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(["Name", "Email", "Company", "Tags", "Source", "Added", "Unsubscribed"])
        for c in db().execute(f"SELECT c.* FROM contacts c WHERE {cond} ORDER BY c.name COLLATE NOCASE", args):
            w.writerow([c["name"], c["email"], c["organization"], c["tags"], c["source"], c["created_at"][:10],
                        (c["unsubscribed_at"] or "")[:10]])
        return Response("﻿" + out.getvalue(), mimetype="text/csv",
                        headers={"Content-Disposition": "attachment; filename=client-list.csv"})

    # ------------------------------------------------ campaigns
    @app.route("/campaigns")
    @A.login_required
    def campaigns_list():
        _resume_stalled()
        cond, args = campaign_scope()
        rows = db().execute(f"SELECT * FROM campaigns WHERE {cond} ORDER BY COALESCE(sent_at, created_at) DESC, id DESC",
                            args).fetchall()
        items = [{**dict(r), **stats(r["id"])} for r in rows]
        return render_template("campaigns.html", campaigns=items, list_count=audience_count("all"))

    @app.route("/campaigns/new", methods=["POST"])
    @A.login_required
    def campaign_new():
        cur = db().execute("INSERT INTO campaigns (company_id, subject, body, created_by_id, created_by, created_at) "
                           "VALUES (?,?,?,?,?,?)",
                           (A.cid(), "", "Hi {first_name},\n\n\n\nThank you,\n" + (g.user["name"] if g.user["id"] else
                                                                                  g.company["name"]),
                            g.user["id"], A.actor(), now()))
        db().commit()
        return redirect(url_for("campaign_edit", camp_id=cur.lastrowid))

    def save_draft(camp_id):
        audience = request.form.get("audience", "all")
        if not (audience == "all" or audience.startswith("tag:")):
            audience = "all"
        url = request.form.get("button_url", "").strip()[:500]
        if url and not re.match(r"^https?://", url):
            url = "https://" + url
        db().execute("UPDATE campaigns SET subject=?, body=?, button_text=?, button_url=?, audience=? WHERE id=?",
                     (request.form.get("subject", "").strip()[:200], request.form.get("body", "")[:20000],
                      request.form.get("button_text", "").strip()[:60], url, audience, camp_id))
        db().commit()

    def problems_before_sending(camp):
        if not camp["subject"].strip():
            return "Add a subject line."
        if len(camp["body"].strip()) < 10:
            return "Write the message first."
        if not A.mailer.enabled():
            return "Email isn't turned on for this server yet (SMTP settings). Ask your FormSign provider."
        if not (g.company.get("mailing_address") or "").strip():
            return ("Add your company's mailing address on the Company page first. Canada's anti-spam law requires it "
                    "in every marketing email." if A.is_manager() else
                    "A manager needs to add the company's mailing address on the Company page first "
                    "(required by Canada's anti-spam law).")
        return None

    @app.route("/campaigns/<int:camp_id>/edit", methods=["GET", "POST"])
    @A.login_required
    def campaign_edit(camp_id):
        camp = load_campaign(camp_id)
        if camp["status"] != "draft":
            return redirect(url_for("campaign_view", camp_id=camp_id))
        if request.method == "POST":
            save_draft(camp_id)
            camp = load_campaign(camp_id)
            action = request.form.get("action", "save")
            if action == "test":
                to = g.user["email"] if g.user["id"] else (g.company.get("notify_email") or "")
                problem = problems_before_sending(camp)
                if problem and "mailing address" not in problem:
                    flash(problem, "error")
                elif not EMAIL_RE.match(to or ""):
                    flash("There's no email address to send the test to.", "error")
                else:
                    try:
                        token = "test-" + secrets.token_urlsafe(8)
                        base = A.BASE_URL or request.url_root.rstrip("/")
                        html, text = render_email(camp, g.company, {"name": g.user["name"], "organization": ""},
                                                  f"{base}/contacts")
                        conn = A.mailer.open_connection()
                        try:
                            conn.send_message(A.mailer.campaign_message(
                                to, "[Test] " + camp["subject"], html, text, g.company["name"],
                                g.company.get("notify_email") or None, f"{base}/contacts"))
                        finally:
                            A.mailer.close_connection(conn)
                        flash(f"Test sent to {to}. Check how it looks before sending to everyone." +
                              (" (Add your mailing address on the Company page before the real send.)"
                               if problem else ""), "ok")
                    except Exception as exc:  # noqa: BLE001
                        flash(f"The test email didn't go out ({exc}).", "error")
            elif action == "review":
                problem = problems_before_sending(camp)
                if problem:
                    flash(problem, "error")
                else:
                    return redirect(url_for("campaign_confirm", camp_id=camp_id))
            else:
                flash("Draft saved.", "ok")
            return redirect(url_for("campaign_edit", camp_id=camp_id))
        counts = {"all": audience_count("all")}
        tags = tags_in_scope()
        for t in tags:
            counts["tag:" + t] = audience_count("tag:" + t)
        return render_template("campaign_edit.html", c=camp, tags=tags, counts=counts)

    @app.route("/campaigns/<int:camp_id>/preview")
    @A.login_required
    def campaign_preview(camp_id):
        camp = load_campaign(camp_id)
        html, _ = render_email(camp, g.company, {"name": "Jordan Sample", "organization": "Sample Co"}, "#")
        return Response(html, mimetype="text/html",
                        headers={"Content-Security-Policy": "default-src 'none'; img-src *; style-src 'unsafe-inline'"})

    @app.route("/campaigns/<int:camp_id>/send", methods=["GET", "POST"])
    @A.login_required
    def campaign_confirm(camp_id):
        camp = load_campaign(camp_id)
        if camp["status"] != "draft":
            return redirect(url_for("campaign_view", camp_id=camp_id))
        problem = problems_before_sending(camp)
        if problem:
            flash(problem, "error")
            return redirect(url_for("campaign_edit", camp_id=camp_id))
        count = audience_count(camp["audience"])
        used = sent_last_24h()
        if request.method == "POST":
            if request.form.get("confirm") != "on":
                flash("Tick the box to confirm everyone on this list agreed to receive your emails.", "error")
            elif count == 0:
                flash("Nobody to send to: this audience has no subscribed clients.", "error")
            elif used + count > DAILY_LIMIT:
                flash(f"This would go over your limit of {DAILY_LIMIT} campaign emails per 24 hours "
                      f"({used} sent already). Send to a smaller group, or try again tomorrow.", "error")
            else:
                sql, args = audience_sql(camp["audience"])
                people = db().execute(f"SELECT c.id, c.email, c.name, c.organization {sql}", args).fetchall()
                for p in people:
                    db().execute("INSERT INTO campaign_recipients (campaign_id, contact_id, email, name, organization, "
                                 "token) VALUES (?,?,?,?,?,?)",
                                 (camp_id, p["id"], p["email"], p["name"], p["organization"], secrets.token_urlsafe(18)))
                db().execute("UPDATE campaigns SET status='sending', sent_at=?, base_url=? WHERE id=?",
                             (now(), A.BASE_URL or request.url_root.rstrip("/"), camp_id))
                db().commit()
                start_sending(camp_id)
                flash(f"Sending to {len(people)} client{'s' if len(people) != 1 else ''}. You can leave this page; "
                      f"it keeps going in the background.", "ok")
                return redirect(url_for("campaign_view", camp_id=camp_id))
        return render_template("campaign_confirm.html", c=camp, count=count, used=used, limit=DAILY_LIMIT,
                               minutes=max(1, round(count / PER_MINUTE)))

    @app.route("/campaigns/<int:camp_id>")
    @A.login_required
    def campaign_view(camp_id):
        camp = load_campaign(camp_id)
        if camp["status"] == "draft":
            return redirect(url_for("campaign_edit", camp_id=camp_id))
        _resume_stalled()
        show = request.args.get("show", "all")
        where = {"failed": " AND status='failed'", "opened": " AND opened_at IS NOT NULL",
                 "queued": " AND status='queued'"}.get(show, "")
        recips = db().execute(f"SELECT * FROM campaign_recipients WHERE campaign_id=?{where} "
                              "ORDER BY CASE status WHEN 'failed' THEN 0 ELSE 1 END, id LIMIT 500", (camp_id,)).fetchall()
        return render_template("campaign_view.html", c=camp, s=stats(camp_id), recips=recips, show=show)

    @app.route("/campaigns/<int:camp_id>/action", methods=["POST"])
    @A.login_required
    def campaign_action(camp_id):
        camp = load_campaign(camp_id)
        action = request.form.get("action")
        if action == "cancel" and camp["status"] == "sending":
            db().execute("UPDATE campaigns SET status='cancelled', finished_at=? WHERE id=?", (now(), camp_id))
            n = db().execute("UPDATE campaign_recipients SET status='skipped', error='cancelled' "
                             "WHERE campaign_id=? AND status='queued'", (camp_id,)).rowcount
            flash(f"Stopped. {n} email{'s' if n != 1 else ''} that hadn't gone out yet won't be sent.", "ok")
        elif action == "retry" and camp["status"] in ("sent", "cancelled"):
            n = db().execute("UPDATE campaign_recipients SET status='queued', error='' WHERE campaign_id=? AND "
                             "status='failed'", (camp_id,)).rowcount
            if n:
                db().execute("UPDATE campaigns SET status='sending', finished_at=NULL WHERE id=?", (camp_id,))
                db().commit()
                start_sending(camp_id)
                flash(f"Retrying {n} email{'s' if n != 1 else ''}.", "ok")
        elif action == "duplicate":
            cur = db().execute("INSERT INTO campaigns (company_id, subject, body, button_text, button_url, audience, "
                               "created_by_id, created_by, created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                               (A.cid(), camp["subject"], camp["body"], camp["button_text"], camp["button_url"],
                                camp["audience"], g.user["id"], A.actor(), now()))
            db().commit()
            flash("Copy created. Edit it and send when ready.", "ok")
            return redirect(url_for("campaign_edit", camp_id=cur.lastrowid))
        elif action == "delete" and camp["status"] == "draft":
            db().execute("DELETE FROM campaigns WHERE id=?", (camp_id,))
            db().commit()
            flash("Draft deleted.", "ok")
            return redirect(url_for("campaigns_list"))
        db().commit()
        return redirect(url_for("campaign_view", camp_id=camp_id))

    # ------------------------------------------------ background sending
    def start_sending(camp_id):
        if app.config.get("CAMPAIGNS_INLINE"):  # tests: send right away in this request
            _send_loop(camp_id)
            return
        threading.Thread(target=_send_loop, args=(camp_id,), daemon=True, name=f"campaign-{camp_id}").start()

    def _claim(conn, camp_id):
        stale = (datetime.now(timezone.utc) - timedelta(seconds=STALE_SECONDS)).replace(microsecond=0).isoformat()
        n = conn.execute("UPDATE campaigns SET worker=?, heartbeat_at=? WHERE id=? AND status='sending' AND "
                         "(worker IS NULL OR worker=? OR heartbeat_at IS NULL OR heartbeat_at < ?)",
                         (WORKER, now(), camp_id, WORKER, stale)).rowcount
        conn.commit()
        return n == 1

    def _send_loop(camp_id):
        """Send the queued emails of one campaign, a few at a time, keeping a heartbeat so that a restarted
        server can pick up where it stopped."""
        with app.app_context():
            conn = db()
            if not _claim(conn, camp_id):
                return
            camp = conn.execute("SELECT * FROM campaigns WHERE id=?", (camp_id,)).fetchone()
            company = A.get_company(camp["company_id"])
            pause = 60.0 / PER_MINUTE
            smtp = None
            try:
                while True:
                    if not _claim(conn, camp_id):
                        return  # cancelled, or another server process took over
                    batch = conn.execute("SELECT * FROM campaign_recipients WHERE campaign_id=? AND status='queued' "
                                         "ORDER BY id LIMIT 20", (camp_id,)).fetchall()
                    if not batch:
                        conn.execute("UPDATE campaigns SET status='sent', finished_at=?, worker=NULL WHERE id=? "
                                     "AND status='sending'", (now(), camp_id))
                        conn.commit()
                        return
                    for r in batch:
                        unsub = conn.execute("SELECT unsubscribed_at FROM contacts WHERE company_id=? AND email=?",
                                             (camp["company_id"], r["email"])).fetchone()
                        if unsub and unsub[0]:
                            conn.execute("UPDATE campaign_recipients SET status='skipped', error='unsubscribed' "
                                         "WHERE id=?", (r["id"],))
                            continue
                        base = camp["base_url"] or ""
                        unsub_url = f"{base}/u/{r['token']}"
                        html, text = render_email(camp, company, dict(r), unsub_url, f"{base}/o/{r['token']}.gif")
                        msg = A.mailer.campaign_message(r["email"], personalize(camp["subject"], dict(r)), html, text,
                                                        company["name"], company.get("notify_email") or None, unsub_url)
                        error = None
                        for attempt in range(2):
                            try:
                                if smtp is None:
                                    smtp = A.mailer.open_connection()
                                smtp.send_message(msg)
                                error = None
                                break
                            except Exception as exc:  # noqa: BLE001  (dropped connection: reconnect once)
                                error = str(exc)[:300] or exc.__class__.__name__
                                A.mailer.close_connection(smtp)
                                smtp = None
                        if error:
                            conn.execute("UPDATE campaign_recipients SET status='failed', error=? WHERE id=?",
                                         (error, r["id"]))
                        else:
                            conn.execute("UPDATE campaign_recipients SET status='sent', error='', sent_at=? WHERE id=?",
                                         (now(), r["id"]))
                        conn.commit()
                        if not app.config.get("CAMPAIGNS_INLINE"):
                            time.sleep(pause)
                    conn.commit()
            except Exception:  # noqa: BLE001
                app.logger.exception("Campaign %s stopped unexpectedly; it will resume on the next page view", camp_id)
            finally:
                A.mailer.close_connection(smtp)

    def _resume_stalled():
        """Restart sending for campaigns whose server process stopped (deploy or restart) mid-send."""
        stale = (datetime.now(timezone.utc) - timedelta(seconds=STALE_SECONDS)).replace(microsecond=0).isoformat()
        for (camp_id,) in db().execute("SELECT id FROM campaigns WHERE company_id=? AND status='sending' AND "
                                       "(heartbeat_at IS NULL OR heartbeat_at < ?)", (A.cid(), stale)).fetchall():
            start_sending(camp_id)
