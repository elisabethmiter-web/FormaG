"""On-page previews of uploaded files, so clients can read a document right above where they sign.

* PDFs are rendered to page images on the server (works on every phone, no browser plug-ins).
* Word, Excel, PowerPoint, OpenDocument, RTF, text and CSV files are converted to PDF with LibreOffice
  when it is installed (the Docker image includes it) and then shown as page images, exactly as laid out.
* Without LibreOffice, Word (.docx), Excel (.xlsx), CSV and text files fall back to a simple HTML
  preview of their content.
* Anything else (ZIP, CAD drawings...) can't be previewed and is offered as a download.
"""
import csv
import os
import shutil
import subprocess
import tempfile
import threading
from html import escape

OFFICE_EXTS = {"doc", "docx", "odt", "rtf", "xls", "xlsx", "ods", "ppt", "pptx", "odp", "csv", "txt"}
_LOCK = threading.Lock()          # one LibreOffice / PDFium job at a time per server process
_PDFIUM_LOCK = threading.Lock()


def ext_of(name):
    return os.path.splitext(name or "")[1].lower().lstrip(".")


def soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def _convert_to_pdf(src, out_pdf):
    exe = soffice()
    if not exe:
        return False
    with _LOCK, tempfile.TemporaryDirectory() as tmp:
        # copy under a plain name: LibreOffice names its output after the input file
        ext = os.path.splitext(src)[1]
        work = os.path.join(tmp, "doc" + ext)
        shutil.copyfile(src, work)
        try:
            subprocess.run([exe, f"-env:UserInstallation=file://{tmp}/profile", "--headless", "--norestore",
                            "--convert-to", "pdf", "--outdir", tmp, work],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120, check=False)
        except (subprocess.TimeoutExpired, OSError):
            return False
        produced = os.path.join(tmp, "doc.pdf")
        if not os.path.exists(produced) or os.path.getsize(produced) == 0:
            return False
        shutil.move(produced, out_pdf)
        return True


# ---------------------------------------------------------------- HTML fallbacks
def _docx_html(path):
    import docx  # python-docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = "".join(
                (f"<strong>{escape(r.text)}</strong>" if r.bold else escape(r.text)) for r in p.runs) or escape(p.text)
            if not text.strip():
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style.startswith("heading") or style == "title":
                out.append(f"<h4>{text}</h4>")
            elif "list" in style:
                out.append(f"<p class='li'>• {text}</p>")
            else:
                out.append(f"<p>{text}</p>")
        elif tag == "tbl":
            t = Table(child, d)
            rows = "".join("<tr>" + "".join(f"<td>{escape(c.text)}</td>" for c in r.cells) + "</tr>"
                           for r in t.rows[:300])
            out.append(f"<div class='table-wrap'><table>{rows}</table></div>")
    return "\n".join(out) or "<p><em>This document has no text to show.</em></p>"


def _rows_html(rows, max_rows=300, max_cols=30):
    body, n = [], 0
    for r in rows:
        if n >= max_rows:
            body.append(f"<tr><td colspan='{max_cols}'><em>… more rows in the file</em></td></tr>")
            break
        cells = ["" if v is None else str(v) for v in list(r)[:max_cols]]
        if not any(c.strip() for c in cells):
            continue
        body.append("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in cells) + "</tr>")
        n += 1
    return "<div class='table-wrap'><table class='sheet'>" + "".join(body) + "</table></div>"


def _xlsx_html(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:5]:
        parts.append(f"<h4>{escape(ws.title)}</h4>" + _rows_html(ws.iter_rows(values_only=True)))
    wb.close()
    return "\n".join(parts)


def _csv_html(path):
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
        return _rows_html(csv.reader(fh))


def _txt_html(path):
    with open(path, encoding="utf-8", errors="replace") as fh:
        text = fh.read(200_000)
    return f"<pre>{escape(text)}</pre>"


def _html_fallback(path, ext):
    if ext == "docx":
        return _docx_html(path)
    if ext == "xlsx":
        return _xlsx_html(path)
    if ext == "csv":
        return _csv_html(path)
    if ext == "txt":
        return _txt_html(path)
    return None


# ---------------------------------------------------------------- public API
def ensure_preview(upload_dir, stored, original_name):
    """Return {"type": "pdf"|"html", "file": stored_name} for a non-PDF upload, or None.
    Results (and failures) are cached next to the upload so the work happens once."""
    ext = ext_of(original_name)
    if ext not in OFFICE_EXTS:
        return None
    stem = os.path.splitext(stored)[0]
    pdf_name, html_name, none_name = f"prev-{stem}.pdf", f"prev-{stem}.html", f"prev-{stem}.none"
    for name, typ in ((pdf_name, "pdf"), (html_name, "html")):
        if os.path.exists(os.path.join(upload_dir, name)):
            return {"type": typ, "file": name}
    if os.path.exists(os.path.join(upload_dir, none_name)):
        return None
    src = os.path.join(upload_dir, stored)
    if _convert_to_pdf(src, os.path.join(upload_dir, pdf_name)):
        return {"type": "pdf", "file": pdf_name}
    try:
        html = _html_fallback(src, ext)
    except Exception:  # noqa: BLE001  (damaged or password-protected file)
        html = None
    if html:
        with open(os.path.join(upload_dir, html_name), "w", encoding="utf-8") as fh:
            fh.write(html)
        return {"type": "html", "file": html_name}
    open(os.path.join(upload_dir, none_name), "w").close()
    return None


def page_count(pdf_path):
    import pypdfium2 as pdfium
    with _PDFIUM_LOCK:
        doc = pdfium.PdfDocument(pdf_path)
        try:
            return len(doc)
        finally:
            doc.close()


def page_png(pdf_path, n, out_png, scale=1.6):
    """Render page n of a PDF to PNG (cached by the caller via out_png)."""
    if os.path.exists(out_png):
        return out_png
    import pypdfium2 as pdfium
    with _PDFIUM_LOCK:
        doc = pdfium.PdfDocument(pdf_path)
        try:
            page = doc[n]
            img = page.render(scale=scale, may_draw_forms=True).to_pil().convert("RGB")
            page.close()
        finally:
            doc.close()
    tmp = out_png + ".tmp"
    img.save(tmp, "PNG", optimize=True)
    os.replace(tmp, out_png)
    return out_png
